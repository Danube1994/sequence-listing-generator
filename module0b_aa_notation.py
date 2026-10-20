"""
模块0b：氨基酸紧凑记法解析与渲染模块

把氨基酸/多肽序列 Word 文档里的连字符记法字符串（形如
    Tyr-Aib-Glu-Gly-Thr-Phe-...-Pro-Ser-NH2
    Tyr-Aib-Glu-...-Lys(AEEA-AEEA-γGlu-Eic)-Ala-...-Pro-Ser-NH2
解析成结构化的逐残基信息，供后续模块（Phase 3）渲染成 ST.26 的
MOD_RES / INSDSeq_sequence。跟核苷酸的模块0（module0_compact_notation.py）
是同一个定位，但语法形状完全不同，业务规则也不一样，所以是独立的
文件，不是同一套代码改出来的。

语法（用 <> 标出的部分可以通过 GUI 里的"氨基酸符号定义"界面配置，
语法结构本身——连字符连接主链、圆括号挂支链、末尾后缀——不可改）：
    (<残基缩写>(\\(<圆括号内容>\\))?)("-"(<残基缩写>(\\(<圆括号内容>\\))?))*
        ("-"<末端后缀符号>)?

- 主链是连字符连接的残基缩写序列。每个缩写要么是 Annex I 表3 的标准
  三字母代码（Tyr/Ala/Glu/...），要么是"氨基酸符号定义"里登记的非
  标准缩写（Aib/Cpa/...），两者都不是就报错。
- 任意一个残基缩写后面可以紧跟一个圆括号，表示这个残基（必须是能
  正常解析出裸序列字母的那种，不能是本身已经是"非标准缩写"的残基，
  这种组合目前没有规则可依，会报错）带了一段接头/偶联支链。圆括号
  内容本身也是连字符连接的片段序列，每个片段要么是标准三字母代码，
  要么在"氨基酸符号定义"的接头/偶联片段名单里；片段前面如果紧跟一
  个希腊字母（中间没有分隔符，比如 "γGlu"），渲染时替换成"对应英文
  单词+空格"（得到 "gamma Glu"）。圆括号内容开头如果多一个连字符
  （比如 "Lys(-Gly-Gly-...)"，这是记法本身的书写习惯，不是每次都会
  有），解析时直接忽略，不当成错误。
- 主链最后可以再跟一个"-<末端后缀符号>"（默认 "NH2"），表示对最后
  一个残基的末端修饰（默认对应 C 端酰胺化）。

裸序列字母的业务规则（已跟用户确认）：非标准缩写各自对应的裸序列
字母存在"氨基酸符号定义"里，不是从化学结构算出来的——除甲基化、
磷酸化这类"侧链化学骨架基本不变、只加个小基团"的修饰保留母核字母
（mAla -> A、pSer -> S）外，其余一律用 'X'（Aib、Cpa、Iva、mPal ->
X）。带圆括号支链的标准残基（比如 Lys(...)）仍然按它自己的标准字母
处理（Lys -> K），因为圆括号语法本身不改变主链残基是谁，只是额外
挂了一段支链说明。

一个残基后面可以跟多个圆括号（比如 "Cys(γGlu)(-CH2-)"）。每个括号按
内容分类：和"氨基酸符号定义"里桥连接定义表某一行"起点写法"一致的是桥
的起点，和"终点写法"一致的是桥的终点，其余是上面说的接头/偶联支链。
起点和终点按出现顺序一一配对（每个终点配它前面最早没配对、终点写法
一致的起点），配成一对就生成一条 CROSSLNK/DISULFID 特征（位置 x..y，
note 按定义表的模板生成，{aa1}/{aa2} 是起点/终点残基的三字母代码）。
桥只能挂在标准三字母残基上；对不上的起点/终点、同一个残基带多个接头
支链，都直接报错。

语法之外的任何残留字符（未登记的缩写、圆括号内容里不认识的片段、
缺分隔符等）一律报错，不静默跳过、不猜测——这套记法是给机器读的，
出现语法之外的情况大概率是数据录入错误，或者是符号定义里还没收录
的新缩写/新片段，需要人工确认。
"""

import re
from dataclasses import dataclass, field

from aa_symbol_legend import load_config
from annex_i_data import AMINO_ACID_SYMBOLS
from module0_compact_notation import NotationParseError

_MAIN_TOKEN_RE = re.compile(r"^([^()\s]+)((?:\([^()]*\))*)$")
_GROUP_RE = re.compile(r"\(([^()]*)\)")


def _standard_code_to_letter():
    return {code: letter for letter, (code, _definition) in AMINO_ACID_SYMBOLS.items()}


def _standard_code_to_definition():
    return {code: definition for _letter, (code, definition) in AMINO_ACID_SYMBOLS.items()}


def _split_top_level(text, sep="-"):
    """
    按分隔符切分主链，但圆括号内部的分隔符不算数（圆括号内容本身也
    用同一个分隔符隔开内部片段，不能被主链切分逻辑误伤）。
    """
    tokens = []
    depth = 0
    current = []
    for ch in text:
        if ch == "(":
            depth += 1
            current.append(ch)
        elif ch == ")":
            depth -= 1
            if depth < 0:
                raise ValueError("圆括号不匹配（多了一个右括号）")
            current.append(ch)
        elif ch == sep and depth == 0:
            tokens.append("".join(current))
            current = []
        else:
            current.append(ch)
    if depth != 0:
        raise ValueError("圆括号不匹配（少了右括号）")
    tokens.append("".join(current))
    return tokens


def _build_bracket_token_re(config):
    """
    按当前符号配置动态生成圆括号内容的 token 正则。片段名（标准三
    字母代码 + 接头片段名单）按长度从长到短排列，保证像 "PEG-2" 这种
    片段名本身带连字符时，不会被更短的片段名提前截断匹配。
    """
    codes = {code for code, _definition in AMINO_ACID_SYMBOLS.values()}
    vocab = sorted(codes | set(config["conjugate_fragments"]), key=len, reverse=True)
    vocab_alt = "|".join(re.escape(v) for v in vocab)

    greek_keys = sorted(config["greek_letter_words"].keys(), key=len, reverse=True)
    if greek_keys:
        greek_alt = "|".join(re.escape(g) for g in greek_keys)
        pattern = (
            rf"(?:(?P<greek>{greek_alt})(?P<greek_term>{vocab_alt}))"
            rf"|(?P<plain_term>{vocab_alt})"
        )
    else:
        pattern = rf"(?P<plain_term>{vocab_alt})"
    return re.compile(pattern)


@dataclass
class DecodedAAResidue:
    position: int
    base_letter: str
    notes: list = field(default_factory=list)  # 这个位置对应的全部 MOD_RES /note 文本


@dataclass
class DecodedAACrosslink:
    start: int
    end: int
    key: str       # CROSSLNK 或 DISULFID
    note: str


@dataclass
class DecodedAASequence:
    row_name: str
    raw_notation: str
    residues: list = field(default_factory=list)
    crosslinks: list = field(default_factory=list)   # List[DecodedAACrosslink]

    @property
    def bare_sequence(self):
        return "".join(r.base_letter for r in self.residues)


def _render_bracket_content(bracket_text, row_name, raw, config, token_re):
    if bracket_text.startswith("-"):
        bracket_text = bracket_text[1:]
    if not bracket_text:
        raise NotationParseError(row_name, "圆括号内容为空", raw)

    rendered = []
    pos = 0
    n = len(bracket_text)
    while pos < n:
        m = token_re.match(bracket_text, pos)
        if not m:
            raise NotationParseError(
                row_name,
                f"圆括号内容位置 {pos} 附近出现无法识别的片段: {bracket_text[pos:]!r}"
                "（可能是新的接头/偶联片段名称，需要先在氨基酸符号定义里添加；"
                "也可能是原始记法录入错误）",
                raw,
            )
        if m.group("greek"):
            word = config["greek_letter_words"][m.group("greek")]
            rendered.append(f"{word} {m.group('greek_term')}")
        else:
            rendered.append(m.group("plain_term"))
        pos = m.end()
        if pos < n:
            if bracket_text[pos] != "-":
                raise NotationParseError(
                    row_name,
                    f"圆括号内容位置 {pos} 附近应该是 '-' 分隔符，实际是 {bracket_text[pos]!r}",
                    raw,
                )
            pos += 1
            if pos >= n:
                raise NotationParseError(row_name, "圆括号内容以多余的 '-' 结尾", raw)

    return "-".join(rendered)


def parse_aa_notation(notation, row_name, config=None):
    """
    解析一条氨基酸紧凑记法字符串，返回 DecodedAASequence。不认识的
    缩写/片段一律报错。config 缺省时从 aa_symbol_legend.load_config()
    读取当前配置。
    """
    config = config if config is not None else load_config()
    code_to_letter = _standard_code_to_letter()
    code_to_definition = _standard_code_to_definition()
    bracket_token_re = _build_bracket_token_re(config)
    abbreviations = config["residue_abbreviations"]

    raw = notation
    text = (notation or "").strip()
    if not text:
        raise NotationParseError(row_name, "记法字符串为空", raw)

    try:
        tokens = _split_top_level(text, "-")
    except ValueError as e:
        raise NotationParseError(row_name, str(e), raw)

    suffix_symbol = config["terminal_suffix_symbol"]
    has_terminal_suffix = bool(tokens) and tokens[-1] == suffix_symbol
    if has_terminal_suffix:
        tokens = tokens[:-1]

    if not tokens:
        raise NotationParseError(row_name, "没有解析出任何残基", raw)

    bridges = config.get("bridges", [])
    opener_map = {b["opener"]: b for b in bridges}
    closer_texts = {b["closer"] for b in bridges}
    pending_openers = []   # 还没配上终点的桥起点
    crosslinks = []

    residues = []
    for i, token in enumerate(tokens):
        position = i + 1
        m = _MAIN_TOKEN_RE.match(token)
        if not m:
            raise NotationParseError(
                row_name,
                f"位置 {position} 的记法 {token!r} 不符合语法"
                "（应该是不含空格和括号的缩写，可选跟若干个圆括号；检查是否多了空格）",
                raw,
            )
        base_part = m.group(1)
        groups = _GROUP_RE.findall(m.group(2))
        note = None

        # 每个括号按内容分类：桥起点 / 桥终点 / 接头支链
        bridge_events = []      # [("open", bridge) 或 ("close", 终点写法)]
        conjugate_groups = []
        for g in groups:
            if g in opener_map:
                bridge_events.append(("open", opener_map[g]))
            elif g in closer_texts:
                bridge_events.append(("close", g))
            else:
                conjugate_groups.append(g)

        if base_part in abbreviations:
            if groups:
                raise NotationParseError(
                    row_name,
                    f"位置 {position} 的 {base_part!r} 既是非标准缩写又带了圆括号，"
                    "这种组合目前没有规则可依，需要人工确认",
                    raw,
                )
            entry = abbreviations[base_part]
            base_letter = entry["bare_letter"]
            note = entry["note"]
        elif base_part in code_to_letter:
            base_letter = code_to_letter[base_part]
            if len(conjugate_groups) > 1:
                raise NotationParseError(
                    row_name,
                    f"位置 {position} 的 {base_part!r} 带了 {len(conjugate_groups)} 个接头/偶联"
                    "支链括号，同一个残基带多个支链目前没有规则可依，需要人工确认",
                    raw,
                )
            if conjugate_groups:
                bracket_note = _render_bracket_content(
                    conjugate_groups[0], row_name, raw, config, bracket_token_re
                )
                note = f"{base_part}({bracket_note})"

            for kind, value in bridge_events:
                if kind == "open":
                    pending_openers.append({"bridge": value, "start": position, "aa": base_part})
                    continue
                idx = next(
                    (k for k, p in enumerate(pending_openers)
                     if p["bridge"]["closer"] == value and p["start"] != position),
                    None,
                )
                if idx is None:
                    raise NotationParseError(
                        row_name,
                        f"位置 {position} 的桥连接终点 ({value}) 前面没有还没配对、"
                        "终点写法一致的起点（终点必须出现在起点后面，也不能跟起点在同一个残基上）",
                        raw,
                    )
                opener = pending_openers.pop(idx)
                try:
                    bridge_note = opener["bridge"]["note"].format(
                        aa1=opener["aa"], aa2=base_part
                    )
                except (KeyError, IndexError, ValueError) as e:
                    raise NotationParseError(
                        row_name,
                        f"桥连接 {opener['bridge']['opener']!r} 的 note 模板无法使用"
                        f"（{e!r}），只支持 {{aa1}}、{{aa2}} 两个占位符",
                        raw,
                    )
                crosslinks.append(DecodedAACrosslink(
                    start=opener["start"], end=position,
                    key=opener["bridge"]["key"], note=bridge_note,
                ))
        else:
            raise NotationParseError(
                row_name,
                f"位置 {i + 1} 的缩写 {base_part!r} 无法识别——既不是 Annex I 表3 的"
                "标准三字母代码，也不在当前氨基酸符号定义里，需要先确认"
                "（可能是新缩写需要添加，也可能是原始记法的拼写错误）",
                raw,
            )

        residues.append(
            DecodedAAResidue(position=i + 1, base_letter=base_letter, notes=[note] if note else [])
        )

    if has_terminal_suffix:
        last_base = _MAIN_TOKEN_RE.match(tokens[-1]).group(1)
        if last_base in abbreviations:
            last_name = abbreviations[last_base]["note"]
        else:
            last_name = code_to_definition[last_base]
        template = config["terminal_suffix_note_template"]
        residues[-1].notes.append(template.format(name=last_name))

    if pending_openers:
        p = pending_openers[0]
        raise NotationParseError(
            row_name,
            f"位置 {p['start']} 的桥连接起点 ({p['bridge']['opener']}) 后面没有找到对应的终点 "
            f"({p['bridge']['closer']})",
            raw,
        )

    return DecodedAASequence(
        row_name=row_name, raw_notation=raw, residues=residues, crosslinks=crosslinks
    )


def load_aa_notation_table(docx_path, table_index=0):
    """
    从氨基酸序列 Word 文件（表格：SEQ ID NO: + 序列 两列）读取，返回
    List[(name, notation)]。SEQ ID 那一列可能带形如
    "20\\n（化合物1）" 这样的附加中文标注，这里只取开头的数字部分
    作为 name；取不到数字就把原始单元格文本整个当作 name，交给上层
    在报错时能看到完整内容。
    """
    import docx

    doc = docx.Document(docx_path)
    if not doc.tables:
        raise NotationParseError("<文件>", f"{docx_path} 里没有找到任何表格")
    table = doc.tables[table_index]
    if len(table.columns) < 2:
        raise NotationParseError(
            "<文件>", f"表格只有 {len(table.columns)} 列，至少需要2列（SEQ ID + 序列）"
        )

    pairs = []
    for row in table.rows[1:]:
        raw_id = row.cells[0].text.strip()
        notation = row.cells[1].text.strip()
        if not raw_id:
            continue
        m = re.match(r"\s*(\d+)", raw_id)
        name = m.group(1) if m else raw_id
        pairs.append((name, notation))
    return pairs


def records_from_aa_docx(docx_path, config=None, table_index=0):
    """
    一步到位：氨基酸 Word 文件 -> (decoded_list, errors)。
    decoded_list: List[DecodedAASequence]，解析成功的行。
    errors: List[NotationParseError]，解析失败的行，不同行互不影响，
        不会中断整体处理。
    """
    config = config if config is not None else load_config()
    pairs = load_aa_notation_table(docx_path, table_index)

    decoded_list = []
    errors = []
    for name, notation in pairs:
        try:
            decoded_list.append(parse_aa_notation(notation, name, config))
        except NotationParseError as e:
            errors.append(e)
    return decoded_list, errors


if __name__ == "__main__":
    import sys

    src = sys.argv[1] if len(sys.argv) > 1 else "SRP260114-序列信息.docx"
    config = load_config()

    decoded_list, errors = records_from_aa_docx(src, config)
    for d in decoded_list:
        mods = [(r.position, n) for r in d.residues for n in r.notes]
        print(f"Name={d.row_name}\t长度={len(d.bare_sequence)}\t{d.bare_sequence}")
        for pos, note in mods:
            print(f"    MOD_RES {pos}: {note}")

    print()
    print(f"共 {len(decoded_list) + len(errors)} 行，成功 {len(decoded_list)} 行，出错 {len(errors)} 行")
    for e in errors:
        print(" -", e)

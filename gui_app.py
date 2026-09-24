"""
ST.26 序列表 XML 生成工具（桌面 GUI）

把核苷酸（模块0-7 + 合并单元格适配器）和氨基酸（模块0b、5b、6b）两条
流水线打包成一个界面小工具：
- 选择序列文件（Excel 或 Word 都行）：
    Excel  -> 自动识别"一行一条序列"和"合并单元格分块"两种排布，
              不需要用户自己判断该用哪个（核苷酸专用格式，氨基酸目前
              只支持 Word 输入）
    Word   -> 先看记法内容自动判断这是核苷酸紧凑记法还是氨基酸连字符
              记法（_is_aa_notation_docx），再分别按"符号定义"/"氨基酸
              符号定义"里配置的规则解析（模块0 / 模块0b）。两条记法
              语法互斥，判断依据见 _is_aa_notation_docx 的说明。
- 点"生成 XML"：跑完整条流水线，校验通过就弹出保存对话框导出 XML；
  有问题就在日志区列出来，并可以导出成 Excel 校验报告

Annex I 参照表（Table 1-4）已经内嵌在 annex_i_data.py 里，不需要每次
再单独选择/上传，界面上也就不再放这个入口。

不重新实现校验/生成逻辑，全部复用已经测试过的模块：
    核苷酸紧凑记法解析 -> module0.records_from_docx（"符号定义"配置）
    核苷酸格式探测/适配 -> adapter_merged_block_reader
    核苷酸校验流水线   -> module7.run_full_validation（内部串了模块1-4，
                          缺省就用内嵌的 Annex I 数据）
    核苷酸 XML 生成    -> module6.build_sequence_listing + module5.escape_apostrophes
    核苷酸错误报告     -> module7.write_report
    氨基酸记法解析+校验 -> module0b.records_from_aa_docx（一步到位，不需要
                          核苷酸模块1-4那条校验链，见 module6b 的说明）
    氨基酸 XML 生成    -> module6b.build_aa_sequence_listing
"""

import os
import re
import sys
import tempfile
import threading
import traceback
import xml.etree.ElementTree as ET
import tkinter as tk
from tkinter import filedialog, messagebox, scrolledtext

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import openpyxl

from adapter_merged_block_reader import load_merged_block_records, write_normalized_xlsx
from module0_compact_notation import records_from_docx, NotationParseError
from module0b_aa_notation import records_from_aa_docx
from module5_xml_generation import build_sequence_data, escape_apostrophes
from module6_batch_driver import build_sequence_listing, XML_DECLARATION, DOCTYPE_LINE
from module6b_aa_batch_driver import build_aa_sequence_listing
from module7_validation_report import run_full_validation, write_report, normalize_issue
from symbol_legend import load_config, save_config, validate_config, LegendError, MAX_ENTRIES, MIN_BLANK_SLOTS
import aa_symbol_legend


def detect_format(xlsx_path):
    """靠 A/B 列是否存在合并单元格判断输入格式：'merged' 或 'flat'。"""
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    for merged in ws.merged_cells.ranges:
        if merged.min_col in (1, 2):
            return "merged"
    return "flat"


def _is_aa_notation_docx(path):
    """
    读表格里第一条非空记法，判断这份 Word 文档该走氨基酸解析器
    （module0b）还是核苷酸解析器（module0）。

    两套记法语法在这一点上是互斥的：核苷酸紧凑记法主链本身不含连字
    符（只有 [接头名] 方括号内部偶尔会带连字符，比如
    "[N3T12-012-01]"）；氨基酸记法主链恰恰是用连字符连接残基缩写的，
    圆括号支链内部也可能带连字符（比如 "PEG-2"）。所以去掉 [...] 和
    (...) 包裹的内容之后，剩下的主链文本里还有没有连字符，就是可靠
    的判断依据——不是猜，是两边语法定型时就已经互斥的结构性差异。
    """
    import docx

    doc = docx.Document(path)
    if not doc.tables:
        return False
    table = doc.tables[0]
    for row in table.rows[1:]:
        if len(row.cells) < 2:
            continue
        notation = row.cells[1].text.strip()
        if not notation:
            continue
        stripped = re.sub(r"\[[^\]]*\]", "", notation)
        stripped = re.sub(r"\([^)]*\)", "", stripped)
        return "-" in stripped
    return False


def run_aa_pipeline(path, log):
    """氨基酸 Word 记法专用流水线：不经过核苷酸模块1-4，模块0b 解析成功即代表校验通过。"""
    config = aa_symbol_legend.load_config()
    decoded_list, notation_errors = records_from_aa_docx(path, config)
    issues = [normalize_issue(e, "模块0b 氨基酸记法解析") for e in notation_errors]

    log("检测到输入格式：Word 氨基酸记法")
    log(f"通过校验 {len(decoded_list)} 条，出错 {len(issues)} 条")

    if not decoded_list:
        return None, issues

    root = build_aa_sequence_listing(decoded_list, file_name="sequence_listing.xml")
    ET.indent(root)
    body = escape_apostrophes(ET.tostring(root, encoding="unicode"))
    xml_text = XML_DECLARATION + "\n" + DOCTYPE_LINE + "\n" + body + "\n"
    return xml_text, issues


def prepare_flat_source(path):
    """
    返回 (flat_xlsx_path, is_temp, fmt, notation_issues)。
    - .docx：走模块0解析紧凑记法，解析失败的行汇总进 notation_issues
      （不中断整体处理），解析成功的行写成临时 flat xlsx。
    - .xlsx 合并单元格分块：先用适配器展开，写一份临时 flat 文件。
    - .xlsx 一行一条序列：直接返回原路径，不用临时文件。
    is_temp=True 时调用方用完要自己删除 flat_xlsx_path。
    """
    ext = os.path.splitext(path)[1].lower()

    if ext == ".docx":
        config = load_config()
        records, notation_errors = records_from_docx(path, config)
        notation_issues = [normalize_issue(e, "模块0 紧凑记法解析") for e in notation_errors]
        fd, tmp_path = tempfile.mkstemp(suffix=".xlsx")
        os.close(fd)
        write_normalized_xlsx(records, tmp_path)
        return tmp_path, True, "word紧凑记法", notation_issues

    fmt = detect_format(path)
    if fmt == "flat":
        return path, False, fmt, []

    records = load_merged_block_records(path)
    fd, tmp_path = tempfile.mkstemp(suffix=".xlsx")
    os.close(fd)
    write_normalized_xlsx(records, tmp_path)
    return tmp_path, True, fmt, []


def run_pipeline(path, log):
    """返回 (xml_text 或 None, issues: List[ValidationIssue])。Annex I 用内嵌数据。"""
    if os.path.splitext(path)[1].lower() == ".docx" and _is_aa_notation_docx(path):
        return run_aa_pipeline(path, log)

    flat_path, is_temp, fmt, notation_issues = prepare_flat_source(path)
    fmt_label = {
        "merged": "合并单元格分块", "flat": "一行一条序列", "word紧凑记法": "Word 紧凑记法",
    }.get(fmt, fmt)
    log(f"检测到输入格式：{fmt_label}")
    if notation_issues:
        log(f"紧凑记法解析失败 {len(notation_issues)} 行（详见校验报告）")
    try:
        ok_records, issues = run_full_validation(flat_path)
    finally:
        if is_temp:
            os.remove(flat_path)

    issues = notation_issues + issues
    log(f"通过校验 {len(ok_records)} 条，出错 {len(issues)} 条")

    if not ok_records:
        return None, issues

    sequence_data_elements = [build_sequence_data(record, features) for record, features in ok_records]
    root = build_sequence_listing(sequence_data_elements, file_name="sequence_listing.xml")
    ET.indent(root)
    body = escape_apostrophes(ET.tostring(root, encoding="unicode"))
    xml_text = XML_DECLARATION + "\n" + DOCTYPE_LINE + "\n" + body + "\n"
    return xml_text, issues


class SymbolLegendDialog(tk.Toplevel):
    """
    "符号定义"界面：模块0紧凑记法解析器里全部符号的可视化编辑器，分两块：

    1. 残基修饰符号表（可以一直加行）：每行两个输入框，符号（比如 m）
       + 说明（比如 2'-O-methoxy modification (2'-OMe)）。
    2. 三个结构性符号各自的定义（各只有一条）：磷硫代酯键符号、端基
       前缀符号、接头说明文字模板。这三个符号在紧凑记法语法里各自的
       "角色"（谁是前缀、谁是残基间连接符、谁包在方括号里）是固定的，
       不能改，但符号本身长什么样、对应什么说明文字，都可以改。

    打开时按已有配置回填；"添加一行"给残基修饰符号表用，到 MAX_ENTRIES
    行封顶；"保存"时过滤掉残基修饰符号表里整行空白的条目，校验通过才
    写盘、关闭窗口，校验不通过就地提示、不关闭。
    """

    def __init__(self, master):
        super().__init__(master)
        self.title("符号定义")
        self.geometry("560x600")
        self.transient(master)
        self.grab_set()

        config = load_config()

        # ---- 残基修饰符号表 ----
        header = tk.Frame(self, padx=10)
        header.pack(fill="x", pady=(10, 0))
        tk.Label(header, text="残基修饰符号", width=18, anchor="w",
                 font=("", 9, "bold")).pack(side="left")
        tk.Label(header, text="对应的 ST.26 修饰说明", anchor="w",
                 font=("", 9, "bold")).pack(side="left", fill="x", expand=True)

        tk.Label(
            self,
            text=f"紧跟在碱基字母后面、描述这个残基本身带什么修饰的符号，最多 {MAX_ENTRIES} 条。",
            fg="#666", padx=10, anchor="w", justify="left",
        ).pack(fill="x")

        # 可滚动区域：Canvas + 内部 Frame + 纵向滚动条，标准 Tkinter 套路
        body = tk.Frame(self)
        body.pack(fill="both", expand=True, padx=10, pady=5)
        canvas = tk.Canvas(body, highlightthickness=0, height=200)
        scrollbar = tk.Scrollbar(body, orient="vertical", command=canvas.yview)
        self.rows_frame = tk.Frame(canvas)
        self.rows_frame.bind(
            "<Configure>", lambda e: canvas.configure(scrollregion=canvas.bbox("all"))
        )
        canvas.create_window((0, 0), window=self.rows_frame, anchor="nw")
        canvas.configure(yscrollcommand=scrollbar.set)
        canvas.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")

        self.row_widgets = []  # List[(symbol_entry, meaning_entry)]

        for symbol, meaning in config["residue_modifiers"].items():
            self._add_row(symbol, meaning)
        while len(self.row_widgets) < MIN_BLANK_SLOTS:
            self._add_row("", "")

        add_row_frm = tk.Frame(self, pady=4)
        add_row_frm.pack()
        self.add_btn = tk.Button(add_row_frm, text="添加一行", width=14, command=self.on_add_row)
        self.add_btn.pack(side="left", padx=5)
        self.count_label = tk.Label(add_row_frm, fg="#666")
        self.count_label.pack(side="left", padx=5)
        self._update_count_label()

        # ---- 三个结构性符号 ----
        struct_frm = tk.LabelFrame(self, text="结构性符号（角色固定，符号和说明可以改）", padx=10, pady=8)
        struct_frm.pack(fill="x", padx=10, pady=(4, 0))
        struct_frm.columnconfigure(1, weight=0)
        struct_frm.columnconfigure(3, weight=1)

        tk.Label(struct_frm, text="磷硫代酯键符号：").grid(row=0, column=0, sticky="w")
        self.ps_symbol_entry = tk.Entry(struct_frm, width=8)
        self.ps_symbol_entry.insert(0, config["ps_link_symbol"])
        self.ps_symbol_entry.grid(row=0, column=1, sticky="w", padx=(0, 10))
        tk.Label(struct_frm, text="说明：").grid(row=0, column=2, sticky="w")
        self.ps_note_entry = tk.Entry(struct_frm)
        self.ps_note_entry.insert(0, config["ps_link_note"])
        self.ps_note_entry.grid(row=0, column=3, sticky="ew")

        tk.Label(struct_frm, text="端基前缀符号：").grid(row=1, column=0, sticky="w", pady=(6, 0))
        self.vp_symbol_entry = tk.Entry(struct_frm, width=8)
        self.vp_symbol_entry.insert(0, config["vp_prefix_symbol"])
        self.vp_symbol_entry.grid(row=1, column=1, sticky="w", padx=(0, 10), pady=(6, 0))
        tk.Label(struct_frm, text="说明：").grid(row=1, column=2, sticky="w", pady=(6, 0))
        self.vp_note_entry = tk.Entry(struct_frm)
        self.vp_note_entry.insert(0, config["vp_note"])
        self.vp_note_entry.grid(row=1, column=3, sticky="ew", pady=(6, 0))

        tk.Label(struct_frm, text="接头说明模板：").grid(row=2, column=0, sticky="w", pady=(6, 0))
        self.linker_template_entry = tk.Entry(struct_frm)
        self.linker_template_entry.insert(0, config["linker_note_template"])
        self.linker_template_entry.grid(row=2, column=1, columnspan=3, sticky="ew", pady=(6, 0))
        tk.Label(
            struct_frm, text="必须包含 {name} 占位符，运行时替换成方括号里实际写的接头名称",
            fg="#666", anchor="w",
        ).grid(row=3, column=0, columnspan=4, sticky="w", pady=(2, 0))

        btn_frm = tk.Frame(self, pady=8)
        btn_frm.pack()
        tk.Button(btn_frm, text="保存", width=14, command=self.on_save).pack(side="left", padx=5)
        tk.Button(btn_frm, text="取消", width=14, command=self.destroy).pack(side="left", padx=5)

    def _update_count_label(self):
        self.count_label.config(text=f"当前 {len(self.row_widgets)} / {MAX_ENTRIES} 行")
        self.add_btn.config(state="disabled" if len(self.row_widgets) >= MAX_ENTRIES else "normal")

    def _add_row(self, symbol="", meaning=""):
        row = tk.Frame(self.rows_frame)
        row.pack(fill="x", pady=1)
        symbol_entry = tk.Entry(row, width=16)
        symbol_entry.insert(0, symbol)
        symbol_entry.pack(side="left")
        meaning_entry = tk.Entry(row)
        meaning_entry.insert(0, meaning)
        meaning_entry.pack(side="left", fill="x", expand=True, padx=(6, 0))
        self.row_widgets.append((symbol_entry, meaning_entry))

    def on_add_row(self):
        if len(self.row_widgets) >= MAX_ENTRIES:
            messagebox.showwarning("提示", f"最多只能有 {MAX_ENTRIES} 条符号定义")
            return
        self._add_row("", "")
        self._update_count_label()

    def on_save(self):
        residue_modifiers = {}
        for symbol_entry, meaning_entry in self.row_widgets:
            symbol = symbol_entry.get().strip()
            meaning = meaning_entry.get().strip()
            if not symbol and not meaning:
                continue  # 整行空白，跳过
            if not symbol or not meaning:
                messagebox.showerror("错误", "符号和说明必须同时填写，不能只填一边")
                return
            if symbol in residue_modifiers:
                messagebox.showerror("错误", f"符号 {symbol!r} 重复定义了")
                return
            residue_modifiers[symbol] = meaning

        if not residue_modifiers:
            messagebox.showerror("错误", "残基修饰符号至少需要一条有效定义")
            return

        config = {
            "residue_modifiers": residue_modifiers,
            "ps_link_symbol": self.ps_symbol_entry.get().strip(),
            "ps_link_note": self.ps_note_entry.get().strip(),
            "vp_prefix_symbol": self.vp_symbol_entry.get().strip(),
            "vp_note": self.vp_note_entry.get().strip(),
            "linker_note_template": self.linker_template_entry.get().strip(),
        }

        try:
            validate_config(config)
            save_config(config)
        except LegendError as e:
            messagebox.showerror("错误", str(e))
            return

        messagebox.showinfo(
            "完成",
            f"已保存 {len(residue_modifiers)} 条残基修饰符号 + 3 个结构性符号的定义",
        )
        self.destroy()


class AASymbolLegendDialog(tk.Toplevel):
    """
    "氨基酸符号定义"界面：氨基酸紧凑记法解析器（模块0b）里全部符号的
    可视化编辑器，对应 aa_symbol_legend.py 的配置结构，分四块：

    1. 残基修饰缩写表（可以一直加行）：每行三个输入框，缩写（比如
       Aib）+ 完整修饰名称/note（比如 2-Aminoisobutyric acid）+
       裸序列字母（比如 X，或者甲基化/磷酸化这类保留母核字母的场合
       填对应字母，比如 A、S）。
    2. 末端后缀符号（只有一条）：符号本身（默认 NH2）+ 说明文字模板
       （必须包含 {name} 占位符，运行时替换成最后一个残基的英文全
       称）。
    3. 希腊字母替换表（可以一直加行）：圆括号子语法里希腊字母 ->
       对应英文单词，比如 γ -> gamma。
    4. 接头/偶联片段名单（可以一直加行，只有一列）：圆括号子语法里
       允许出现的、本身不是氨基酸的片段名称，比如 AEEA、Eic、PEG-2。

    整个对话框内容放进一个可滚动区域，保存/取消按钮固定在底部不随
    内容滚动。保存时校验通过才写盘、关闭窗口，不通过就地提示、不
    关闭，逻辑跟核苷酸版 SymbolLegendDialog 一致。
    """

    def __init__(self, master):
        super().__init__(master)
        self.title("氨基酸符号定义")
        self.geometry("680x700")
        self.transient(master)
        self.grab_set()

        config = aa_symbol_legend.load_config()

        outer = tk.Frame(self)
        outer.pack(fill="both", expand=True)
        canvas = tk.Canvas(outer, highlightthickness=0)
        scrollbar = tk.Scrollbar(outer, orient="vertical", command=canvas.yview)
        content = tk.Frame(canvas)
        content.bind("<Configure>", lambda e: canvas.configure(scrollregion=canvas.bbox("all")))
        canvas.create_window((0, 0), window=content, anchor="nw")
        canvas.configure(yscrollcommand=scrollbar.set)
        canvas.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")

        # ---- 1. 残基修饰缩写表 ----
        header = tk.Frame(content, padx=10)
        header.pack(fill="x", pady=(10, 0))
        tk.Label(header, text="缩写", width=12, anchor="w", font=("", 9, "bold")).pack(side="left")
        tk.Label(header, text="完整修饰名称（note）", width=32, anchor="w",
                 font=("", 9, "bold")).pack(side="left")
        tk.Label(header, text="裸序列字母", anchor="w", font=("", 9, "bold")).pack(side="left")

        tk.Label(
            content,
            text=f"主链里不在 Annex I 表3标准20种氨基酸符号里的缩写，最多 {aa_symbol_legend.MAX_ENTRIES} 条。"
                 "裸序列字母填 'X'，或者甲基化/磷酸化这类保留母核字母场合对应的标准单字母符号。",
            fg="#666", padx=10, anchor="w", justify="left", wraplength=640,
        ).pack(fill="x")

        self.residue_rows_frame = tk.Frame(content)
        self.residue_rows_frame.pack(fill="x", padx=10, pady=5)
        self.residue_rows = []  # List[(abbr_entry, note_entry, bare_letter_entry)]

        for abbr, entry in config["residue_abbreviations"].items():
            self._add_residue_row(abbr, entry["note"], entry["bare_letter"])
        while len(self.residue_rows) < aa_symbol_legend.MIN_BLANK_SLOTS:
            self._add_residue_row("", "", "")

        residue_btn_frm = tk.Frame(content, pady=4)
        residue_btn_frm.pack()
        self.residue_add_btn = tk.Button(residue_btn_frm, text="添加一行", width=14,
                                          command=self.on_add_residue_row)
        self.residue_add_btn.pack(side="left", padx=5)
        self.residue_count_label = tk.Label(residue_btn_frm, fg="#666")
        self.residue_count_label.pack(side="left", padx=5)
        self._update_residue_count_label()

        # ---- 2. 末端后缀符号 ----
        suffix_frm = tk.LabelFrame(content, text="末端后缀符号（角色固定：跟在主链最后一个残基后面）",
                                    padx=10, pady=8)
        suffix_frm.pack(fill="x", padx=10, pady=(8, 0))
        suffix_frm.columnconfigure(1, weight=0)
        suffix_frm.columnconfigure(3, weight=1)

        tk.Label(suffix_frm, text="符号：").grid(row=0, column=0, sticky="w")
        self.suffix_symbol_entry = tk.Entry(suffix_frm, width=8)
        self.suffix_symbol_entry.insert(0, config["terminal_suffix_symbol"])
        self.suffix_symbol_entry.grid(row=0, column=1, sticky="w", padx=(0, 10))
        tk.Label(suffix_frm, text="说明模板：").grid(row=0, column=2, sticky="w")
        self.suffix_template_entry = tk.Entry(suffix_frm)
        self.suffix_template_entry.insert(0, config["terminal_suffix_note_template"])
        self.suffix_template_entry.grid(row=0, column=3, sticky="ew")
        tk.Label(
            suffix_frm, text="必须包含 {name} 占位符，运行时替换成最后一个残基的英文全称",
            fg="#666", anchor="w",
        ).grid(row=1, column=0, columnspan=4, sticky="w", pady=(2, 0))

        # ---- 3. 希腊字母替换表 ----
        tk.Label(content, text="希腊字母替换表", anchor="w", font=("", 9, "bold"),
                 padx=10).pack(fill="x", pady=(10, 0))
        tk.Label(
            content,
            text="圆括号子语法里紧跟片段名前面、中间没有分隔符的希腊字母，替换成\"英文单词+空格\"。",
            fg="#666", padx=10, anchor="w", justify="left", wraplength=640,
        ).pack(fill="x")

        self.greek_rows_frame = tk.Frame(content)
        self.greek_rows_frame.pack(fill="x", padx=10, pady=5)
        self.greek_rows = []  # List[(greek_entry, word_entry)]

        for greek, word in config["greek_letter_words"].items():
            self._add_greek_row(greek, word)
        while len(self.greek_rows) < aa_symbol_legend.GREEK_MIN_BLANK_SLOTS:
            self._add_greek_row("", "")

        greek_btn_frm = tk.Frame(content, pady=4)
        greek_btn_frm.pack()
        self.greek_add_btn = tk.Button(greek_btn_frm, text="添加一行", width=14,
                                        command=self.on_add_greek_row)
        self.greek_add_btn.pack(side="left", padx=5)
        self.greek_count_label = tk.Label(greek_btn_frm, fg="#666")
        self.greek_count_label.pack(side="left", padx=5)
        self._update_greek_count_label()

        # ---- 4. 接头/偶联片段名单 ----
        tk.Label(content, text="接头/偶联片段名单", anchor="w", font=("", 9, "bold"),
                 padx=10).pack(fill="x", pady=(10, 0))
        tk.Label(
            content,
            text="圆括号内容按连字符切分后，不是标准三字母代码的片段必须在这张名单里，否则解析报错。",
            fg="#666", padx=10, anchor="w", justify="left", wraplength=640,
        ).pack(fill="x")

        self.fragment_rows_frame = tk.Frame(content)
        self.fragment_rows_frame.pack(fill="x", padx=10, pady=5)
        self.fragment_rows = []  # List[fragment_entry]

        for frag in config["conjugate_fragments"]:
            self._add_fragment_row(frag)
        while len(self.fragment_rows) < aa_symbol_legend.FRAGMENT_MIN_BLANK_SLOTS:
            self._add_fragment_row("")

        fragment_btn_frm = tk.Frame(content, pady=4)
        fragment_btn_frm.pack()
        self.fragment_add_btn = tk.Button(fragment_btn_frm, text="添加一行", width=14,
                                           command=self.on_add_fragment_row)
        self.fragment_add_btn.pack(side="left", padx=5)
        self.fragment_count_label = tk.Label(fragment_btn_frm, fg="#666")
        self.fragment_count_label.pack(side="left", padx=5)
        self._update_fragment_count_label()

        # ---- 底部按钮（不随内容滚动）----
        btn_frm = tk.Frame(self, pady=8)
        btn_frm.pack(side="bottom")
        tk.Button(btn_frm, text="保存", width=14, command=self.on_save).pack(side="left", padx=5)
        tk.Button(btn_frm, text="取消", width=14, command=self.destroy).pack(side="left", padx=5)

    # ---- 残基修饰缩写表 ----

    def _add_residue_row(self, abbr="", note="", bare_letter=""):
        row = tk.Frame(self.residue_rows_frame)
        row.pack(fill="x", pady=1)
        abbr_entry = tk.Entry(row, width=12)
        abbr_entry.insert(0, abbr)
        abbr_entry.pack(side="left")
        note_entry = tk.Entry(row, width=34)
        note_entry.insert(0, note)
        note_entry.pack(side="left", padx=(4, 0))
        bare_entry = tk.Entry(row, width=8)
        bare_entry.insert(0, bare_letter)
        bare_entry.pack(side="left", padx=(4, 0))
        self.residue_rows.append((abbr_entry, note_entry, bare_entry))

    def _update_residue_count_label(self):
        self.residue_count_label.config(text=f"当前 {len(self.residue_rows)} / {aa_symbol_legend.MAX_ENTRIES} 行")
        self.residue_add_btn.config(
            state="disabled" if len(self.residue_rows) >= aa_symbol_legend.MAX_ENTRIES else "normal"
        )

    def on_add_residue_row(self):
        if len(self.residue_rows) >= aa_symbol_legend.MAX_ENTRIES:
            messagebox.showwarning("提示", f"最多只能有 {aa_symbol_legend.MAX_ENTRIES} 条缩写定义")
            return
        self._add_residue_row()
        self._update_residue_count_label()

    # ---- 希腊字母替换表 ----

    def _add_greek_row(self, greek="", word=""):
        row = tk.Frame(self.greek_rows_frame)
        row.pack(fill="x", pady=1)
        greek_entry = tk.Entry(row, width=12)
        greek_entry.insert(0, greek)
        greek_entry.pack(side="left")
        word_entry = tk.Entry(row)
        word_entry.insert(0, word)
        word_entry.pack(side="left", fill="x", expand=True, padx=(6, 0))
        self.greek_rows.append((greek_entry, word_entry))

    def _update_greek_count_label(self):
        self.greek_count_label.config(text=f"当前 {len(self.greek_rows)} / {aa_symbol_legend.MAX_ENTRIES} 行")
        self.greek_add_btn.config(
            state="disabled" if len(self.greek_rows) >= aa_symbol_legend.MAX_ENTRIES else "normal"
        )

    def on_add_greek_row(self):
        if len(self.greek_rows) >= aa_symbol_legend.MAX_ENTRIES:
            messagebox.showwarning("提示", f"最多只能有 {aa_symbol_legend.MAX_ENTRIES} 条希腊字母定义")
            return
        self._add_greek_row()
        self._update_greek_count_label()

    # ---- 接头/偶联片段名单 ----

    def _add_fragment_row(self, fragment=""):
        row = tk.Frame(self.fragment_rows_frame)
        row.pack(fill="x", pady=1)
        fragment_entry = tk.Entry(row)
        fragment_entry.insert(0, fragment)
        fragment_entry.pack(side="left", fill="x", expand=True)
        self.fragment_rows.append(fragment_entry)

    def _update_fragment_count_label(self):
        self.fragment_count_label.config(text=f"当前 {len(self.fragment_rows)} / {aa_symbol_legend.MAX_ENTRIES} 行")
        self.fragment_add_btn.config(
            state="disabled" if len(self.fragment_rows) >= aa_symbol_legend.MAX_ENTRIES else "normal"
        )

    def on_add_fragment_row(self):
        if len(self.fragment_rows) >= aa_symbol_legend.MAX_ENTRIES:
            messagebox.showwarning("提示", f"最多只能有 {aa_symbol_legend.MAX_ENTRIES} 条片段名定义")
            return
        self._add_fragment_row()
        self._update_fragment_count_label()

    # ---- 保存 ----

    def on_save(self):
        residue_abbreviations = {}
        for abbr_entry, note_entry, bare_entry in self.residue_rows:
            abbr = abbr_entry.get().strip()
            note = note_entry.get().strip()
            bare_letter = bare_entry.get().strip().upper()
            if not abbr and not note and not bare_letter:
                continue  # 整行空白，跳过
            if not abbr or not note or not bare_letter:
                messagebox.showerror("错误", "缩写、完整修饰名称、裸序列字母必须同时填写，不能只填一部分")
                return
            if abbr in residue_abbreviations:
                messagebox.showerror("错误", f"缩写 {abbr!r} 重复定义了")
                return
            residue_abbreviations[abbr] = {"note": note, "bare_letter": bare_letter}

        if not residue_abbreviations:
            messagebox.showerror("错误", "残基修饰缩写至少需要一条有效定义")
            return

        greek_letter_words = {}
        for greek_entry, word_entry in self.greek_rows:
            greek = greek_entry.get().strip()
            word = word_entry.get().strip()
            if not greek and not word:
                continue
            if not greek or not word:
                messagebox.showerror("错误", "希腊字母和对应英文单词必须同时填写，不能只填一边")
                return
            if greek in greek_letter_words:
                messagebox.showerror("错误", f"希腊字母 {greek!r} 重复定义了")
                return
            greek_letter_words[greek] = word

        conjugate_fragments = []
        seen_fragments = set()
        for fragment_entry in self.fragment_rows:
            frag = fragment_entry.get().strip()
            if not frag:
                continue
            if frag in seen_fragments:
                messagebox.showerror("错误", f"接头/偶联片段名称 {frag!r} 重复了")
                return
            seen_fragments.add(frag)
            conjugate_fragments.append(frag)

        config = {
            "residue_abbreviations": residue_abbreviations,
            "terminal_suffix_symbol": self.suffix_symbol_entry.get().strip(),
            "terminal_suffix_note_template": self.suffix_template_entry.get().strip(),
            "greek_letter_words": greek_letter_words,
            "conjugate_fragments": conjugate_fragments,
        }

        try:
            aa_symbol_legend.validate_config(config)
            aa_symbol_legend.save_config(config)
        except aa_symbol_legend.LegendError as e:
            messagebox.showerror("错误", str(e))
            return

        messagebox.showinfo(
            "完成",
            f"已保存 {len(residue_abbreviations)} 条残基修饰缩写 + 末端后缀符号 + "
            f"{len(greek_letter_words)} 条希腊字母替换 + {len(conjugate_fragments)} 条接头片段名单",
        )
        self.destroy()


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("ST.26 序列表 XML 生成工具")
        self.geometry("700x500")
        self.minsize(600, 400)

        self.source_path = tk.StringVar()
        self.last_issues = []
        self.pending_xml = None

        top = tk.Frame(self, padx=10, pady=10)
        top.pack(fill="x")
        top.columnconfigure(1, weight=1)

        tk.Label(top, text="序列文件（Excel 或 Word）：").grid(row=0, column=0, sticky="w")
        tk.Entry(top, textvariable=self.source_path).grid(row=0, column=1, sticky="ew", padx=5)
        tk.Button(top, text="浏览...", command=self.pick_source).grid(row=0, column=2)

        btn_frm = tk.Frame(self, pady=6)
        btn_frm.pack()
        self.gen_btn = tk.Button(btn_frm, text="生成 XML", width=20, command=self.on_generate)
        self.gen_btn.pack(side="left", padx=5)
        self.report_btn = tk.Button(btn_frm, text="导出校验报告", width=20,
                                     command=self.on_export_report, state="disabled")
        self.report_btn.pack(side="left", padx=5)
        tk.Button(btn_frm, text="符号定义...", width=20,
                  command=self.on_edit_symbol_legend).pack(side="left", padx=5)
        tk.Button(btn_frm, text="氨基酸符号定义...", width=20,
                  command=self.on_edit_aa_symbol_legend).pack(side="left", padx=5)

        self.log_area = scrolledtext.ScrolledText(self, width=80, height=20)
        self.log_area.pack(fill="both", expand=True, padx=10, pady=(0, 10))

    def log(self, msg):
        self.log_area.insert("end", str(msg) + "\n")
        self.log_area.see("end")
        self.update_idletasks()

    def pick_source(self):
        path = filedialog.askopenfilename(
            title="选择序列文件",
            filetypes=[("Excel / Word", "*.xlsx *.docx"), ("Excel", "*.xlsx"), ("Word", "*.docx")],
        )
        if path:
            self.source_path.set(path)

    def on_generate(self):
        src = self.source_path.get().strip()
        if not src:
            messagebox.showwarning("提示", "请先选择序列文件（Excel 或 Word）")
            return
        if not os.path.isfile(src):
            messagebox.showerror("错误", "文件路径不存在，请重新选择")
            return
        if os.path.splitext(src)[1].lower() not in (".xlsx", ".docx"):
            messagebox.showerror("错误", "只支持 .xlsx 或 .docx 文件")
            return

        self.gen_btn.config(state="disabled")
        self.report_btn.config(state="disabled")
        self.log_area.delete("1.0", "end")
        threading.Thread(target=self._generate_worker, args=(src,), daemon=True).start()

    def _generate_worker(self, src):
        try:
            xml_text, issues = run_pipeline(src, self.log)
        except Exception:
            self.log("发生异常，未生成 XML：")
            self.log(traceback.format_exc())
            self.after(0, lambda: self.gen_btn.config(state="normal"))
            return

        self.last_issues = issues
        if issues:
            self.after(0, lambda: self.report_btn.config(state="normal"))

        if xml_text is None:
            self.log("全部行都未通过校验，没有可用数据，未生成 XML。请导出校验报告修正 Excel 后重试。")
            self.after(0, lambda: self.gen_btn.config(state="normal"))
        else:
            self.pending_xml = xml_text
            self.after(0, self._save_xml_dialog)

    def _save_xml_dialog(self):
        path = filedialog.asksaveasfilename(
            title="保存 XML 文件", defaultextension=".xml",
            filetypes=[("XML", "*.xml")], initialfile="sequence_listing.xml",
        )
        if path:
            with open(path, "w", encoding="utf-8") as f:
                f.write(self.pending_xml)
            self.log(f"已保存: {path}")
            if self.last_issues:
                self.log(f"注意：仍有 {len(self.last_issues)} 条记录因校验失败被跳过，"
                         "未包含在这份 XML 里，建议导出校验报告修正后重新生成。")
            else:
                messagebox.showinfo("完成", f"XML 已生成：{path}")
        else:
            self.log("已取消保存")
        self.gen_btn.config(state="normal")

    def on_edit_symbol_legend(self):
        SymbolLegendDialog(self)

    def on_edit_aa_symbol_legend(self):
        AASymbolLegendDialog(self)

    def on_export_report(self):
        if not self.last_issues:
            return
        path = filedialog.asksaveasfilename(
            title="保存校验报告", defaultextension=".xlsx",
            filetypes=[("Excel", "*.xlsx")], initialfile="校验报告.xlsx",
        )
        if not path:
            return
        write_report(self.last_issues, path)
        self.log(f"已保存校验报告: {path}")


def main():
    if "--selftest" in sys.argv:
        idx = sys.argv.index("--selftest")
        src = sys.argv[idx + 1]
        xml_text, issues = run_pipeline(src, print)
        print(f"selftest: xml_generated={xml_text is not None}, issues={len(issues)}")
        sys.exit(1 if xml_text is None else 0)

    app = App()
    app.mainloop()


if __name__ == "__main__":
    main()

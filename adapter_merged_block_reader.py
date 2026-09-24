"""
适配器：把 Excel 里"每个序列一个区块"的各种排布方式统一读出来，转换
成模块1认识的"一行一条序列"格式（Name / Sequence / modification_text
三列，modification_text 是该序列所有 Feature 片段拼接后的字符串）。

同一个 Sheet 里可能混着两种区块形状（测试例-1-14.xlsx 就是这样，
序列1-9、13-14用第一种，序列10-12用第二种，混在同一份文件里）：
    (a) 合并单元格分块：SEQ ID No. 和 sequence 两列做了真正的 Excel
        合并单元格，跨越该序列对应的所有行；Feature 列每行一条，不合并。
    (b) 单行整段：SEQ ID No./sequence/Feature 三列都只占一行，Feature
        列里已经是分号拼接好的完整字符串（没有合并单元格，因为合并单
        元格至少要跨两行，单行序列不会产生合并记录）。

早期版本只认 (a)——直接扫 ws.merged_cells.ranges 找区块，凡是没有
合并单元格的序列整行整行地从结果里消失，没有任何报错，这正是
测试例-1-14.xlsx 里序列10-12被吃掉的原因：这三条序列各自只占一行，
根本没有生成合并单元格记录，扫合并单元格自然找不到它们。

现在改成按行扫描：A 列出现非空值，就是一个新区块的开始；这一行如果
恰好是某个合并单元格区域的左上角，区块就跟着合并范围延伸到底部（对应
形状 a），否则区块就是这一行本身（对应形状 b）。两种形状统一用同一套
逻辑处理，不需要事先判断整份文件是哪种格式，也不会因为一份文件混着
两种形状而漏行。

拼接 modification_text 时用空字符串直接连接——每条 Feature 值本身已经
带了结尾的 ";"，直接拼接产生的形如 '...";misc_feature...' 正好是模块2
已经验证过能正确处理的"分号后无空格"分隔风格；形状(b)本身就是一整段
已经拼好的文本，拼接列表里只有一个元素，效果等价于原样返回。
"""

import openpyxl


def load_merged_block_records(path, sheet_name=None):
    """返回 List[(name, sequence, modification_text)]，按行首次出现顺序排列。"""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet_name] if sheet_name else wb[wb.sheetnames[0]]

    # A 列"起始行 -> 合并终止行"的映射；不在这个映射里的起始行，
    # 说明它是形状(b)的单行区块，终止行就是它自己。
    col_a_merge_end = {}
    for merged in ws.merged_cells.ranges:
        if merged.min_col == 1:
            col_a_merge_end[merged.min_row] = merged.max_row

    records = []
    for row in ws.iter_rows(min_row=2):
        name_cell = row[0]
        if name_cell.value is None:
            continue  # 上一个区块的续行，已经在处理那个区块时读过了

        min_row = name_cell.row
        max_row = col_a_merge_end.get(min_row, min_row)

        name = ws.cell(min_row, 1).value
        sequence = ws.cell(min_row, 2).value
        feature_parts = []
        for r in range(min_row, max_row + 1):
            val = ws.cell(r, 3).value
            if val is not None and str(val).strip():
                feature_parts.append(str(val).strip())
        modification_text = "".join(feature_parts)
        records.append((name, sequence, modification_text))

    return records


def write_normalized_xlsx(records, out_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "normalized"
    ws.append(["Name", "Sequence", "modification_text"])
    for name, sequence, modification_text in records:
        ws.append([name, sequence, modification_text])
    wb.save(out_path)


if __name__ == "__main__":
    import sys

    src = sys.argv[1] if len(sys.argv) > 1 else "测试例-1-7.xlsx"
    out = sys.argv[2] if len(sys.argv) > 2 else "测试例-1-7_normalized.xlsx"

    records = load_merged_block_records(src)
    print(f"读到 {len(records)} 条序列区块")
    for name, sequence, mod_text in records:
        print(f"  Name={name}\tSequence长度={len(sequence) if sequence else 0}\t"
              f"feature片段数={mod_text.count('misc_feature') + mod_text.count('modified_base')}")

    write_normalized_xlsx(records, out)
    print(f"已写出规范化文件: {out}")

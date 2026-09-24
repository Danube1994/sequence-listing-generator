"""
模块0：紧凑记法解析与渲染模块

把 doc1（原始 word 表1）里的紧凑记法字符串（形如
    UmsGmsGmsAmsUmCmCfUmUfUfUfGmUmAmGmAmUmCmUmUmsGms[N3T12-012-01]
    VPUmsAfsAmGmAmUfCmUmAmCmAmAmAmAfGmGfAmUmCmCmAmsUmsUm
）解析成结构化的逐残基修饰事件（0b），再按项目约定的规则渲染成
现有流水线（模块1-7）认识的 flat modification_text（0c）。

语法（用 <> 标出的四类符号全部可以通过 GUI 里的"符号定义"界面配置，
默认值就是下面示例里的样子，符号本身和含义都可以改，语法结构——谁是
前缀、谁是后缀、方括号包裹接头名——不可改）：
    <vp_prefix_symbol>? ([ACGU]<residue_modifier>?<ps_link_symbol>?)+ (\\[linker名\\])?

- 可选开头前缀符号（默认 "VP"）：1号残基带对应的端基修饰
- 每个残基：一个碱基字母(ACGU) + 可选一个"残基修饰符号"（默认表里有
  m=2'-OMe、f=2'-F，可以增删改） + 可选一个磷硫代酯键符号（默认 "s"，
  表示跟下一个残基或末尾接头之间是这种连接）
- 可选结尾 "[接头名]"：3'端接了一个命名接头分子，说明文字按可配置的
  模板生成

所有符号定义（残基修饰符号表 + 磷硫代酯键符号 + 端基前缀符号 + 接头
说明模板）从 symbol_legend.load_config() 读取，对应 GUI 里的"符号
定义"界面。

渲染规则（用户已确认）：
    碱基内修饰（糖环修饰、碱基修饰、端基前缀修饰等） -> modified_base，
        携带 /mod_base（属于内嵌 Annex I Table 2 列表的修饰用相应缩写，
        否则用 "OTHER"）+ /note="..." 自由文本
    末端 linker（3'/5'端）与序列内部磷硫代酯连接 -> misc_feature，
        只带 /note="..."，不带 /mod_base

语法之外的任何残留字符（多余空格、不认识的符号等）一律报错，不静默
跳过、不猜测——紧凑记法本身是给机器读的，出现语法之外的字符大概率
是数据录入错误，或者是符号定义里还没收录的新符号，需要人工确认，
参见 doc1 第13行 "...Ums Um s[N3T12-012-01]" 那次真实的空格录入错误。
"""

import re
from dataclasses import dataclass, field

from symbol_legend import load_config


class NotationParseError(Exception):
    def __init__(self, row_name, message, raw=""):
        self.row_name = row_name
        self.message = message
        self.raw = raw
        super().__init__(f"[Name={row_name!r}] {message} | 原始记法: {raw!r}")


# ---- 0b：紧凑记法解析 ----

_LINKER_RE = re.compile(r"\[([^\]]*)\]\s*$")


def _build_token_re(config):
    """
    按当前符号配置动态生成 token 正则。残基修饰符号按长度从长到短
    排列，保证多字符符号不会被更短的符号提前截断匹配。
    """
    symbols = sorted(config["residue_modifiers"].keys(), key=len, reverse=True)
    if symbols:
        symbol_group = "(?:" + "|".join(re.escape(s) for s in symbols) + ")"
    else:
        symbol_group = "(?!)"  # 空表：这个分支永远不匹配
    ps_symbol = re.escape(config["ps_link_symbol"])
    return re.compile(r"([ACGU])(" + symbol_group + r")?(" + ps_symbol + r")?")


@dataclass
class DecodedResidue:
    position: int
    base: str
    sugar_mod: str = None        # 残基修饰符号（比如 'm'/'f'）或 None
    ps_link_to_next: bool = False
    vp: bool = False


@dataclass
class DecodedSequence:
    row_name: str
    raw_sequence: str
    residues: list = field(default_factory=list)   # List[DecodedResidue]
    linker_name: str = None                          # 3'端接头名称，没有则 None


def parse_compact_notation(notation, row_name, config=None):
    """
    解析一条紧凑记法字符串，返回 DecodedSequence。不认识的字符一律报错。
    config 缺省时从 symbol_legend.load_config() 读取当前配置。
    """
    config = config if config is not None else load_config()
    token_re = _build_token_re(config)
    vp_prefix = config["vp_prefix_symbol"]

    text = (notation or "").strip()
    if not text:
        raise NotationParseError(row_name, "记法字符串为空")

    vp = False
    if text.startswith(vp_prefix):
        vp = True
        text = text[len(vp_prefix):]

    linker_name = None
    m_linker = _LINKER_RE.search(text)
    if m_linker:
        linker_name = m_linker.group(1)
        text = text[: m_linker.start()]

    residues = []
    pos = 0
    for i, m in enumerate(token_re.finditer(text)):
        if m.start() != pos:
            raise NotationParseError(
                row_name,
                f"位置 {pos} 附近出现无法识别的字符: {text[pos:m.start()]!r}"
                "（如果是新符号，需要先在符号定义里添加）",
                notation,
            )
        base, sugar, s = m.group(1), m.group(2), m.group(3)
        residues.append(DecodedResidue(
            position=i + 1,
            base=base,
            sugar_mod=sugar or None,
            ps_link_to_next=bool(s),
            vp=(vp and i == 0),
        ))
        pos = m.end()

    if pos != len(text):
        raise NotationParseError(
            row_name,
            f"记法末尾出现无法识别的字符: {text[pos:]!r}"
            "（如果是新符号，需要先在符号定义里添加）",
            notation,
        )
    if not residues:
        raise NotationParseError(row_name, "没有解析出任何残基", notation)

    if linker_name is not None and not residues[-1].ps_link_to_next:
        raise NotationParseError(
            row_name,
            f"检测到接头名 {linker_name!r}，但最后一个残基后面没有磷硫代酯键符号"
            f"（{config['ps_link_symbol']!r}），无法确定接头是怎么连接的，"
            "不能默默丢掉这个接头信息",
            notation,
        )

    return DecodedSequence(
        row_name=row_name,
        raw_sequence="".join(r.base for r in residues),
        residues=residues,
        linker_name=linker_name,
    )


# ---- 0c：渲染成 flat modification_text ----

# 修饰说明文字 -> Annex I Table 2 精确缩写（如果有的话）。目前配置表里
# 的修饰都没有跟碱基无关的精确 Table 2 缩写可用，真实数据也统一走
# mod_base="OTHER"（参见对话记录），这里先留空，以后遇到有精确对应
# 关系的修饰再加。
NOTATION_TO_MOD_BASE_ABBREV = {}


def _mod_base_value(note_text):
    return NOTATION_TO_MOD_BASE_ABBREV.get(note_text, "OTHER")


def _modified_base_clause(position, note_text):
    return f'modified_base {position} /mod_base="{_mod_base_value(note_text)}" /note="{note_text}";'


def _misc_feature_clause(location, note_text):
    return f'misc_feature {location} /note="{note_text}";'


def render_modification_text(decoded, config=None):
    """DecodedSequence -> flat modification_text 字符串。config 缺省同上。"""
    config = config if config is not None else load_config()
    residue_modifiers = config["residue_modifiers"]
    clauses = []
    n = len(decoded.residues)

    for r in decoded.residues:
        if r.vp:
            clauses.append(_modified_base_clause(r.position, config["vp_note"]))
        if r.sugar_mod:
            note_text = residue_modifiers.get(r.sugar_mod)
            if note_text is None:
                # 解析阶段能匹配到这个符号，说明当时的 config 里有它；
                # 渲染阶段用的 config 如果跟解析时不一致才会走到这里，
                # 属于调用方用法问题，报错而不是猜。
                raise NotationParseError(
                    decoded.row_name,
                    f"符号 {r.sugar_mod!r} 在当前符号定义里没有对应的说明文字",
                )
            clauses.append(_modified_base_clause(r.position, note_text))
        if r.ps_link_to_next:
            if r.position == n:
                if decoded.linker_name:
                    linker_note = config["linker_note_template"].format(name=decoded.linker_name)
                    clauses.append(_misc_feature_clause(r.position, linker_note))
                else:
                    clauses.append(_misc_feature_clause(
                        f"{r.position}^{r.position + 1}", config["ps_link_note"]
                    ))
            else:
                clauses.append(_misc_feature_clause(
                    f"{r.position}^{r.position + 1}", config["ps_link_note"]
                ))

    return "".join(clauses)


def compact_notation_to_record(name, notation, config=None):
    """一步到位：紧凑记法字符串 -> (name, raw_sequence, modification_text)。"""
    config = config if config is not None else load_config()
    decoded = parse_compact_notation(notation, name, config)
    return name, decoded.raw_sequence, render_modification_text(decoded, config)


def load_compact_notation_table(docx_path, table_index=0):
    """
    从 doc1 格式的 Word 文件（表1：SEQ ID + 紧凑记法两列）读取，返回
    List[(name, notation)]，原样返回不做解析，解析交给 records_from_docx。
    """
    import docx

    doc = docx.Document(docx_path)
    if not doc.tables:
        raise NotationParseError("<文件>", f"{docx_path} 里没有找到任何表格")
    table = doc.tables[table_index]
    if len(table.columns) < 2:
        raise NotationParseError(
            "<文件>", f"表格只有 {len(table.columns)} 列，至少需要2列（SEQ ID + 紧凑记法）"
        )

    pairs = []
    for row in table.rows[1:]:
        name = row.cells[0].text.strip()
        notation = row.cells[1].text.strip()
        if not name:
            continue
        pairs.append((name, notation))
    return pairs


def records_from_docx(docx_path, config=None, table_index=0):
    """
    一步到位：doc1 格式 Word 文件 -> (records, errors)。
    records: List[(name, raw_sequence, modification_text)]，解析成功的行，
        顺序、结构直接对得上 adapter_merged_block_reader 的输出，可以
        原样传给 write_normalized_xlsx。
    errors: List[NotationParseError]，解析失败的行，不会中断整体处理，
        不同行互不影响。
    """
    config = config if config is not None else load_config()
    pairs = load_compact_notation_table(docx_path, table_index)

    records = []
    errors = []
    for name, notation in pairs:
        try:
            records.append(compact_notation_to_record(name, notation, config))
        except NotationParseError as e:
            errors.append(e)
    return records, errors


if __name__ == "__main__":
    import sys
    sys.path.insert(0, ".")

    src = sys.argv[1] if len(sys.argv) > 1 else "1-测试例-序列信息.docx"
    config = load_config()
    print(f"当前符号配置: {config}")
    print()

    records, errors = records_from_docx(src, config)
    for name, seq, mod_text in records:
        print(f"Name={name}\tSequence长度={len(seq)}\t{seq}")

    print()
    print(f"共 {len(records) + len(errors)} 行，成功 {len(records)} 行，出错 {len(errors)} 行")
    for e in errors:
        print(" -", e)
